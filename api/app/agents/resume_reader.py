"""Reads the user's base resume (PDF or DOCX) into ResumeData."""
from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader

from ..llm import LLMError, document_block, structured
from ..schemas import ResumeData

SYSTEM = (
    "You convert resumes into structured data. Copy content faithfully: keep every employer, "
    "title, date, degree and bullet exactly as written. Do not summarize, improve or invent "
    "anything. Use empty strings for fields that are absent. Dates stay in the resume's own format."
)


def extract_text(data: bytes, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines).strip()
    raise ValueError("Upload a PDF or DOCX file.")


def parse_resume(data: bytes, suffix: str, text: str) -> ResumeData:
    if suffix.lower() == ".pdf":
        content = [document_block(data), {"type": "text", "text": "Extract this resume."}]
    else:
        if not text:
            raise LLMError("The document has no readable text.")
        content = [{"type": "text", "text": f"Extract this resume:\n\n{text}"}]
    return structured(SYSTEM, content, ResumeData, effort="low")
